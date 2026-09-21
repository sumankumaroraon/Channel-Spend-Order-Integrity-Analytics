"""
build_excel.py
Builds excel/Channel_Spend_Order_Integrity_Analytics.xlsx — the Excel/Power-Query-style
deliverable for the project. All KPI and summary cells are live formulas against the raw
data tabs (SUMIFS/COUNTIFS/AVERAGE, no hardcoded results), so the workbook recalculates
if the underlying data changes. The Audit Flags tab is an imported pipeline output
(from python/audit_engine.py or sql/audit_flags.sql) — that one is data, not a formula
layer, the same way a real BI workbook would pull in an upstream exceptions feed.
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

ROOT = os.path.join(os.path.dirname(__file__), "..")
RAW_DIR = os.path.join(ROOT, "data", "raw")
OUT_DIR = os.path.join(ROOT, "outputs")
XLSX_DIR = os.path.join(ROOT, "excel")
os.makedirs(XLSX_DIR, exist_ok=True)
XLSX_PATH = os.path.join(XLSX_DIR, "Channel_Spend_Order_Integrity_Analytics.xlsx")

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color="1F3864")
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
LABEL_FONT = Font(name=FONT_NAME, bold=True, size=10)
VALUE_FONT = Font(name=FONT_NAME, size=11, bold=True, color="1F3864")
BODY_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

orders = pd.read_csv(os.path.join(RAW_DIR, "orders.csv"))
spend = pd.read_csv(os.path.join(RAW_DIR, "marketing_spend.csv"))
campaigns = pd.read_csv(os.path.join(RAW_DIR, "campaigns.csv"))
customers = pd.read_csv(os.path.join(RAW_DIR, "customers.csv"))
audit_flags = pd.read_csv(os.path.join(OUT_DIR, "audit_flags.csv"))

wb = Workbook()
wb.remove(wb.active)


def style_header_row(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def write_data_table(ws, df, start_row, table_name, date_cols=None, bool_cols=None, col_widths=None):
    """Writes a dataframe as a native Excel Table starting at start_row, row 1 col 1."""
    date_cols = date_cols or []
    bool_cols = bool_cols or []
    headers = list(df.columns)
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    style_header_row(ws, start_row, len(headers))

    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, (h, val) in enumerate(zip(headers, row), start=1):
            cell = ws.cell(row=i, column=j)
            if pd.isna(val):
                cell.value = None
            elif h in bool_cols:
                cell.value = bool(val)
            elif h in date_cols:
                cell.value = pd.Timestamp(val).date()
                cell.number_format = "yyyy-mm-dd"
            else:
                cell.value = val
            cell.font = BODY_FONT

    end_row = start_row + len(df)
    end_col_letter = get_column_letter(len(headers))
    ref = f"A{start_row}:{end_col_letter}{end_row}"
    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(tbl)

    if col_widths:
        for j, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    else:
        for j in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 16

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return end_row


# ============================================================================
# 1. Read Me
# ============================================================================
ws = wb.create_sheet("Read Me")
ws.sheet_view.showGridLines = False
ws["B2"] = "Channel Spend & Order Integrity Analytics"
ws["B2"].font = TITLE_FONT
ws["B3"] = "Marketing spend, order revenue, and a rule-based audit engine for a simulated e-commerce retailer"
ws["B3"].font = SUBTITLE_FONT

readme_lines = [
    ("", ""),
    ("What this is", ""),
    ("", "A self-contained analytics workbook for a fictional e-commerce retailer, "
         "\u201cNorthfield Retail Co.\u201d All data is synthetically generated "
         "(python/generate_data.py, fixed random seed) for portfolio/demo purposes."),
    ("", ""),
    ("How the numbers work", ""),
    ("", "KPI Summary, Channel Performance, and Monthly Trend are all live formulas "
         "(SUMIFS/COUNTIFS/AVERAGE) against the raw data tabs \u2014 they recalculate if "
         "you edit the raw data. The Audit Flags tab is an imported output from the "
         "companion Python (python/audit_engine.py) and SQL (sql/audit_flags.sql) audit "
         "engines, which agree with each other exactly \u2014 that's the point of building it twice."),
    ("", ""),
    ("Tabs", ""),
    ("", "KPI Summary \u2014 executive KPI cards"),
    ("", "Channel Performance \u2014 revenue, spend, ROAS, refund rate by channel"),
    ("", "Monthly Trend \u2014 revenue vs. spend, Jan\u2013Dec 2025"),
    ("", "Audit Summary \u2014 exception counts and dollars by flag type"),
    ("", "Audit Flags \u2014 the full flagged-record detail (imported)"),
    ("", "Orders / Marketing Spend / Campaigns / Customers \u2014 raw data"),
    ("", ""),
    ("Key assumptions", ""),
    ("", "Discount policy ceiling: 25% (policy_rules in sql/schema.sql)"),
    ("", "Outlier fence: channel Q3 + 4.5x IQR on order_amount"),
    ("", "Chargeback risk flag: 2+ chargebacks from one customer"),
    ("", "Repeated-refund flag: 3+ identical-amount refunds from one customer"),
    ("", "\"Estimated recoverable savings\" = flagged dollars from the five "
         "policy/avoidable-spend flags only (excessive discounts, duplicate charges, "
         "orphaned campaign spend, wasted spend, budget overruns) \u2014 the data-quality-only "
         "flags (orphan IDs, misattribution) are excluded since they aren't recoverable dollars."),
]
r = 5
for label, text in readme_lines:
    if label:
        ws.cell(row=r, column=2, value=label).font = LABEL_FONT
    elif text:
        cell = ws.cell(row=r, column=2, value=text)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 28
    r += 1
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 105

# ============================================================================
# 2. Orders (raw)
# ============================================================================
ws_orders = wb.create_sheet("Orders")
ws_orders.sheet_view.showGridLines = False
orders_out = orders.copy()
end_row_orders = write_data_table(
    ws_orders, orders_out, start_row=1, table_name="OrdersTbl",
    date_cols=["order_date", "refund_date"], bool_cols=["refund_flag", "chargeback_flag"],
)
n_order_cols = len(orders_out.columns)
# helper columns: net_revenue, month, refund_flag_num, chargeback_flag_num
col_map = {c: i + 1 for i, c in enumerate(orders_out.columns)}
net_rev_col = n_order_cols + 1
month_col = n_order_cols + 2
refund_num_col = n_order_cols + 3
cb_num_col = n_order_cols + 4
ws_orders.cell(row=1, column=net_rev_col, value="net_revenue")
ws_orders.cell(row=1, column=month_col, value="month")
ws_orders.cell(row=1, column=refund_num_col, value="refund_flag_num")
ws_orders.cell(row=1, column=cb_num_col, value="chargeback_flag_num")
for c in (net_rev_col, month_col, refund_num_col, cb_num_col):
    cell = ws_orders.cell(row=1, column=c)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.border = BORDER

amt_L = get_column_letter(col_map["order_amount"])
disc_L = get_column_letter(col_map["discount_amount"])
date_L = get_column_letter(col_map["order_date"])
refund_L = get_column_letter(col_map["refund_flag"])
cb_L = get_column_letter(col_map["chargeback_flag"])

for i in range(2, end_row_orders + 1):
    ws_orders.cell(row=i, column=net_rev_col, value=f"={amt_L}{i}-{disc_L}{i}").font = BODY_FONT
    ws_orders.cell(row=i, column=month_col, value=f'=TEXT({date_L}{i},"YYYY-MM")').font = BODY_FONT
    ws_orders.cell(row=i, column=refund_num_col, value=f"=IF({refund_L}{i},1,0)").font = BODY_FONT
    ws_orders.cell(row=i, column=cb_num_col, value=f"=IF({cb_L}{i},1,0)").font = BODY_FONT

for c in (net_rev_col, month_col, refund_num_col, cb_num_col):
    ws_orders.column_dimensions[get_column_letter(c)].width = 14

ORDERS_LAST = end_row_orders
CH_L = get_column_letter(col_map["channel"])
ORDER_ID_L = get_column_letter(col_map["order_id"])
CAMP_L = get_column_letter(col_map["campaign_id"])

# ============================================================================
# 3. Marketing Spend (raw)
# ============================================================================
ws_spend = wb.create_sheet("Marketing Spend")
ws_spend.sheet_view.showGridLines = False
spend_out = spend.copy()
end_row_spend = write_data_table(
    ws_spend, spend_out, start_row=1, table_name="SpendTbl", date_cols=["date"]
)
n_spend_cols = len(spend_out.columns)
spend_col_map = {c: i + 1 for i, c in enumerate(spend_out.columns)}
sp_month_col = n_spend_cols + 1
ws_spend.cell(row=1, column=sp_month_col, value="month")
c0 = ws_spend.cell(row=1, column=sp_month_col)
c0.fill, c0.font, c0.border = HEADER_FILL, HEADER_FONT, BORDER
sp_date_L = get_column_letter(spend_col_map["date"])
for i in range(2, end_row_spend + 1):
    ws_spend.cell(row=i, column=sp_month_col, value=f'=TEXT({sp_date_L}{i},"YYYY-MM")').font = BODY_FONT
ws_spend.column_dimensions[get_column_letter(sp_month_col)].width = 14

SPEND_LAST = end_row_spend
SP_CH_L = get_column_letter(spend_col_map["channel"])
SP_AMT_L = get_column_letter(spend_col_map["spend_amount"])
SP_CAMP_L = get_column_letter(spend_col_map["campaign_id"])

# ============================================================================
# 4. Campaigns (raw + actual spend / variance formulas)
# ============================================================================
ws_camp = wb.create_sheet("Campaigns")
ws_camp.sheet_view.showGridLines = False
camp_out = campaigns.copy()
end_row_camp = write_data_table(
    ws_camp, camp_out, start_row=1, table_name="CampaignsTbl", date_cols=["start_date", "end_date"]
)
camp_col_map = {c: i + 1 for i, c in enumerate(camp_out.columns)}
n_camp_cols = len(camp_out.columns)
actual_col = n_camp_cols + 1
var_col = n_camp_cols + 2
over_col = n_camp_cols + 3
for c, name in [(actual_col, "actual_spend"), (var_col, "variance"), (over_col, "over_budget")]:
    ws_camp.cell(row=1, column=c, value=name)
    cell = ws_camp.cell(row=1, column=c)
    cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER

CAMP_ID_L = get_column_letter(camp_col_map["campaign_id"])
BUDGET_L = get_column_letter(camp_col_map["approved_budget"])
for i in range(2, end_row_camp + 1):
    ws_camp.cell(row=i, column=actual_col,
                 value=f"=SUMIF('Marketing Spend'!{SP_CAMP_L}2:{SP_CAMP_L}{SPEND_LAST},{CAMP_ID_L}{i},"
                       f"'Marketing Spend'!{SP_AMT_L}2:{SP_AMT_L}{SPEND_LAST})").font = BODY_FONT
    actual_L = get_column_letter(actual_col)
    ws_camp.cell(row=i, column=var_col, value=f"={actual_L}{i}-{BUDGET_L}{i}").font = BODY_FONT
    var_L = get_column_letter(var_col)
    ws_camp.cell(row=i, column=over_col, value=f'=IF({var_L}{i}>0,"Over","Within")').font = BODY_FONT
for c in (actual_col, var_col, over_col):
    ws_camp.column_dimensions[get_column_letter(c)].width = 14

CAMP_LAST = end_row_camp
CAMP_ACTUAL_L = get_column_letter(actual_col)
CAMP_OVER_L = get_column_letter(over_col)

# ============================================================================
# 5. Customers (raw)
# ============================================================================
ws_cust = wb.create_sheet("Customers")
ws_cust.sheet_view.showGridLines = False
end_row_cust = write_data_table(
    ws_cust, customers, start_row=1, table_name="CustomersTbl", date_cols=["signup_date"]
)
CUST_LAST = end_row_cust

# ============================================================================
# 6. Audit Flags (imported pipeline output)
# ============================================================================
ws_flags = wb.create_sheet("Audit Flags")
ws_flags.sheet_view.showGridLines = False
ws_flags["A1"] = "Imported from outputs/audit_flags.csv \u2014 generated by python/audit_engine.py (cross-checked against sql/audit_flags.sql)"
ws_flags["A1"].font = SUBTITLE_FONT
end_row_flags = write_data_table(ws_flags, audit_flags, start_row=3, table_name="AuditFlagsTbl")
FLAGS_LAST = end_row_flags
FLAG_COL_L = "C"   # flag
AMT_COL_L = "E"    # amount
# (columns are record_type, record_id, flag, detail, amount -> A..E)

# ============================================================================
# 7. Channel Performance
# ============================================================================
ws_chan = wb.create_sheet("Channel Performance", 2)
ws_chan.sheet_view.showGridLines = False
ws_chan["B2"] = "Channel Performance"
ws_chan["B2"].font = TITLE_FONT
ws_chan["B3"] = "Revenue, spend, orders, refund rate, and ROAS \u2014 all SUMIFS/COUNTIFS formulas against Orders and Marketing Spend"
ws_chan["B3"].font = SUBTITLE_FONT

headers = ["Channel", "Revenue", "Spend", "Orders", "Refund Rate", "ROAS"]
hr = 5
for j, h in enumerate(headers, start=2):
    ws_chan.cell(row=hr, column=j, value=h)
style_header_row(ws_chan, hr, len(headers), start_col=2)

channels = sorted(orders["channel"].unique().tolist())
for i, ch in enumerate(channels, start=hr + 1):
    ws_chan.cell(row=i, column=2, value=ch).font = BODY_FONT
    ws_chan.cell(row=i, column=3,
                 value=f"=SUMIF(Orders!{CH_L}2:{CH_L}{ORDERS_LAST},B{i},"
                       f"Orders!{get_column_letter(net_rev_col)}2:{get_column_letter(net_rev_col)}{ORDERS_LAST})"
                 ).font = BODY_FONT
    ws_chan.cell(row=i, column=4,
                 value=f"=SUMIF('Marketing Spend'!{SP_CH_L}2:{SP_CH_L}{SPEND_LAST},B{i},"
                       f"'Marketing Spend'!{SP_AMT_L}2:{SP_AMT_L}{SPEND_LAST})"
                 ).font = BODY_FONT
    ws_chan.cell(row=i, column=5,
                 value=f"=COUNTIF(Orders!{CH_L}2:{CH_L}{ORDERS_LAST},B{i})").font = BODY_FONT
    ws_chan.cell(row=i, column=6,
                 value=f"=SUMIF(Orders!{CH_L}2:{CH_L}{ORDERS_LAST},B{i},"
                       f"Orders!{get_column_letter(refund_num_col)}2:{get_column_letter(refund_num_col)}{ORDERS_LAST})/E{i}"
                 ).font = BODY_FONT
    ws_chan.cell(row=i, column=7,
                 value=f'=IF(D{i}>0,C{i}/D{i},"N/A")').font = BODY_FONT
    for col, fmt in [(3, '$#,##0'), (4, '$#,##0'), (5, '#,##0'), (6, '0.0%'), (7, '0.00"x"')]:
        ws_chan.cell(row=i, column=col).number_format = fmt
    for c in range(2, 8):
        ws_chan.cell(row=i, column=c).border = BORDER

last_chan_row = hr + len(channels)
for c, w in zip(range(1, 8), [3, 18, 16, 16, 12, 14, 10]):
    ws_chan.column_dimensions[get_column_letter(c)].width = w

bar = BarChart()
bar.title = "Revenue vs. Marketing Spend by Channel"
bar.y_axis.title = "USD"
bar.style = 10
data = Reference(ws_chan, min_col=3, max_col=4, min_row=hr, max_row=last_chan_row)
cats = Reference(ws_chan, min_col=2, min_row=hr + 1, max_row=last_chan_row)
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.height, bar.width = 9, 20
ws_chan.add_chart(bar, f"B{last_chan_row + 3}")

# ============================================================================
# 8. Monthly Trend
# ============================================================================
ws_month = wb.create_sheet("Monthly Trend", 3)
ws_month.sheet_view.showGridLines = False
ws_month["B2"] = "Monthly Trend"
ws_month["B2"].font = TITLE_FONT
ws_month["B3"] = "Revenue vs. marketing spend, Jan\u2013Dec 2025 \u2014 SUMIFS formulas keyed to a TEXT(date,\"YYYY-MM\") helper column"
ws_month["B3"].font = SUBTITLE_FONT

months = [f"2025-{m:02d}" for m in range(1, 13)]
mh = 5
for j, h in enumerate(["Month", "Revenue", "Marketing Spend"], start=2):
    ws_month.cell(row=mh, column=j, value=h)
style_header_row(ws_month, mh, 3, start_col=2)

month_col_L = get_column_letter(month_col)
sp_month_col_L = get_column_letter(sp_month_col)
for i, mo in enumerate(months, start=mh + 1):
    ws_month.cell(row=i, column=2, value=mo).font = BODY_FONT
    ws_month.cell(row=i, column=3,
                  value=f"=SUMIF(Orders!{month_col_L}2:{month_col_L}{ORDERS_LAST},B{i},"
                        f"Orders!{get_column_letter(net_rev_col)}2:{get_column_letter(net_rev_col)}{ORDERS_LAST})"
                  ).font = BODY_FONT
    ws_month.cell(row=i, column=4,
                  value=f"=SUMIF('Marketing Spend'!{sp_month_col_L}2:{sp_month_col_L}{SPEND_LAST},B{i},"
                        f"'Marketing Spend'!{SP_AMT_L}2:{SP_AMT_L}{SPEND_LAST})"
                  ).font = BODY_FONT
    for col in (3, 4):
        ws_month.cell(row=i, column=col).number_format = '$#,##0'
        ws_month.cell(row=i, column=col).border = BORDER
    ws_month.cell(row=i, column=2).border = BORDER

last_month_row = mh + len(months)
for c, w in zip(range(1, 5), [3, 12, 16, 18]):
    ws_month.column_dimensions[get_column_letter(c)].width = w

line = LineChart()
line.title = "Monthly Revenue vs. Marketing Spend \u2014 2025"
line.y_axis.title = "USD"
line.style = 12
data = Reference(ws_month, min_col=3, max_col=4, min_row=mh, max_row=last_month_row)
cats = Reference(ws_month, min_col=2, min_row=mh + 1, max_row=last_month_row)
line.add_data(data, titles_from_data=True)
line.set_categories(cats)
line.height, line.width = 9, 22
ws_month.add_chart(line, f"B{last_month_row + 3}")

# ============================================================================
# 9. Audit Summary
# ============================================================================
ws_asum = wb.create_sheet("Audit Summary", 4)
ws_asum.sheet_view.showGridLines = False
ws_asum["B2"] = "Audit Summary"
ws_asum["B2"].font = TITLE_FONT
ws_asum["B3"] = "Exception counts and dollars by flag type \u2014 COUNTIF/SUMIF formulas against the Audit Flags tab"
ws_asum["B3"].font = SUBTITLE_FONT

flag_types = sorted(audit_flags["flag"].unique().tolist(),
                     key=lambda f: -audit_flags[audit_flags.flag == f].shape[0])
ah = 5
for j, h in enumerate(["Flag Type", "Count", "Total $ Impact"], start=2):
    ws_asum.cell(row=ah, column=j, value=h)
style_header_row(ws_asum, ah, 3, start_col=2)

for i, flag in enumerate(flag_types, start=ah + 1):
    ws_asum.cell(row=i, column=2, value=flag.replace("_", " ").title()).font = BODY_FONT
    ws_asum.cell(row=i, column=3,
                 value=f"=COUNTIF('Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},"
                       f'"{flag}")').font = BODY_FONT
    ws_asum.cell(row=i, column=4,
                 value=f"=SUMIF('Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},"
                       f"\"{flag}\",'Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST})"
                 ).font = BODY_FONT
    ws_asum.cell(row=i, column=4).number_format = '$#,##0'
    for c in range(2, 5):
        ws_asum.cell(row=i, column=c).border = BORDER

last_asum_row = ah + len(flag_types)
total_row = last_asum_row + 1
ws_asum.cell(row=total_row, column=2, value="Total").font = LABEL_FONT
ws_asum.cell(row=total_row, column=3, value=f"=SUM(C{ah+1}:C{last_asum_row})").font = LABEL_FONT
ws_asum.cell(row=total_row, column=4, value=f"=SUM(D{ah+1}:D{last_asum_row})").font = LABEL_FONT
ws_asum.cell(row=total_row, column=4).number_format = '$#,##0'

for c, w in zip(range(1, 5), [3, 30, 10, 16]):
    ws_asum.column_dimensions[get_column_letter(c)].width = w

bar2 = BarChart()
bar2.title = "Flagged Records by Exception Type"
bar2.style = 11
bar2.type = "bar"
data = Reference(ws_asum, min_col=3, max_col=3, min_row=ah, max_row=last_asum_row)
cats = Reference(ws_asum, min_col=2, min_row=ah + 1, max_row=last_asum_row)
bar2.add_data(data, titles_from_data=True)
bar2.set_categories(cats)
bar2.height, bar2.width = 10, 20
ws_asum.add_chart(bar2, f"B{total_row + 3}")

# ============================================================================
# 10. KPI Summary (executive front page)
# ============================================================================
ws_kpi = wb.create_sheet("KPI Summary", 1)
ws_kpi.sheet_view.showGridLines = False
ws_kpi["B2"] = "Executive KPI Summary"
ws_kpi["B2"].font = TITLE_FONT
ws_kpi["B3"] = "Northfield Retail Co. \u2014 FY2025 \u2014 all figures are live formulas against the raw data tabs"
ws_kpi["B3"].font = SUBTITLE_FONT

kpi_rows = [
    ("Total Revenue", f"=SUM(Orders!{get_column_letter(net_rev_col)}2:{get_column_letter(net_rev_col)}{ORDERS_LAST})", '$#,##0'),
    ("Total Marketing Spend", f"=SUM('Marketing Spend'!{SP_AMT_L}2:{SP_AMT_L}{SPEND_LAST})", '$#,##0'),
    ("Blended ROAS", "=C9/C10", '0.00"x"'),
    ("Total Orders", f"=COUNTA(Orders!{ORDER_ID_L}2:{ORDER_ID_L}{ORDERS_LAST})", '#,##0'),
    ("Average Order Value", f"=AVERAGE(Orders!{amt_L}2:{amt_L}{ORDERS_LAST})", '$#,##0.00'),
    ("Refund Rate", f"=SUM(Orders!{get_column_letter(refund_num_col)}2:{get_column_letter(refund_num_col)}{ORDERS_LAST})/C12", '0.0%'),
    ("Chargeback Rate", f"=SUM(Orders!{get_column_letter(cb_num_col)}2:{get_column_letter(cb_num_col)}{ORDERS_LAST})/C12", '0.0%'),
    ("Discount Leakage ($)", f"=SUM(Orders!{disc_L}2:{disc_L}{ORDERS_LAST})", '$#,##0'),
    ("Blended CAC", f"=C10/COUNTIF(Customers!B2:B{CUST_LAST},\">=2025-01-01\")", '$#,##0.00'),
    ("Campaigns Over Budget", f'=COUNTIF(Campaigns!{CAMP_OVER_L}2:{CAMP_OVER_L}{CAMP_LAST},"Over")', '#,##0'),
    ("Budget Compliance Rate", f'=1-(C18/COUNTA(Campaigns!{CAMP_ID_L}2:{CAMP_ID_L}{CAMP_LAST}))', '0.0%'),
    ("Total Flagged Records", f"='Audit Summary'!C{total_row}", '#,##0'),
    ("Total Flagged $ Impact", f"='Audit Summary'!D{total_row}", '$#,##0'),
]
r = 9
for label, formula, fmt in kpi_rows:
    ws_kpi.cell(row=r, column=2, value=label).font = LABEL_FONT
    vcell = ws_kpi.cell(row=r, column=3, value=formula)
    vcell.font = VALUE_FONT
    vcell.number_format = fmt
    for c in (2, 3):
        ws_kpi.cell(row=r, column=c).border = BORDER
    r += 1

ws_kpi["C9"].comment = Comment(
    "Formula: SUM of Orders.net_revenue (order_amount - discount_amount). "
    "net_revenue is itself a formula column in Orders.", "audit_engine"
)
ws_kpi["C17"].comment = Comment(
    'Blended CAC = Total Marketing Spend / new customers with signup_date in 2025. '
    'Simplification: attributes all 2025 spend to 2025 signups (single-year dataset, no prior-year cohort).',
    "audit_engine"
)
ws_kpi["B24"].value = "Estimated Recoverable Savings"
ws_kpi["B24"].font = LABEL_FONT
ws_kpi["C24"] = (f"=SUMIFS('Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST},"
                  f"'Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},\"excessive_discount\")"
                  f"+SUMIFS('Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST},"
                  f"'Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},\"duplicate_order\")"
                  f"+SUMIFS('Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST},"
                  f"'Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},\"orphan_campaign_spend\")"
                  f"+SUMIFS('Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST},"
                  f"'Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},\"wasted_spend_no_orders\")"
                  f"+SUMIFS('Audit Flags'!{AMT_COL_L}4:{AMT_COL_L}{FLAGS_LAST},"
                  f"'Audit Flags'!{FLAG_COL_L}4:{FLAG_COL_L}{FLAGS_LAST},\"budget_overrun\")")
ws_kpi["C24"].font = VALUE_FONT
ws_kpi["C24"].number_format = '$#,##0'
ws_kpi["C24"].comment = Comment(
    "Sum of flagged-$ amount for the five avoidable-spend flag types only "
    "(see Read Me tab). Data-quality-only flags are excluded.", "audit_engine"
)
for c in (2, 3):
    ws_kpi.cell(row=24, column=c).border = BORDER

ws_kpi.column_dimensions["A"].width = 3
ws_kpi.column_dimensions["B"].width = 28
ws_kpi.column_dimensions["C"].width = 18

# ============================================================================
# Sheet order + save
# ============================================================================
order = ["Read Me", "KPI Summary", "Channel Performance", "Monthly Trend", "Audit Summary",
         "Audit Flags", "Orders", "Marketing Spend", "Campaigns", "Customers"]
wb._sheets = [wb[name] for name in order]
wb.active = 0

wb.save(XLSX_PATH)
print("Workbook written:", os.path.abspath(XLSX_PATH))
